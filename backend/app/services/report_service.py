import io
import json
from typing import Dict, Any, Optional
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, PageBreak, KeepTogether,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.schemas import FullAnalysisResponse


BLACK = colors.HexColor("#0a0a0a")
NEON_GREEN = colors.HexColor("#00ff88")
NEON_BLUE = colors.HexColor("#00d4ff")
NEON_PURPLE = colors.HexColor("#bf5af2")
DARK_GRAY = colors.HexColor("#1a1a2e")
MID_GRAY = colors.HexColor("#16213e")
LIGHT_GRAY = colors.HexColor("#e0e0e0")
WHITE = colors.white
RED = colors.HexColor("#ff3b3b")
YELLOW = colors.HexColor("#ffd60a")


def generate_pdf_report(analysis: FullAnalysisResponse) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title="Competitor Intelligence Report — Practical EduSkills",
        author="Jarvis AI War Room",
    )

    styles = _build_styles()
    story = []

    # ── Cover ──────────────────────────────────────────────────────────────────
    story.append(Spacer(1, 1 * cm))
    story.append(Paragraph("🛡️ COMPETITOR INTELLIGENCE REPORT", styles["cover_title"]))
    story.append(Paragraph("Practical EduSkills — Founder War Room", styles["cover_sub"]))
    story.append(Paragraph(f"Generated: {datetime.utcnow().strftime('%d %B %Y, %H:%M UTC')}", styles["cover_date"]))
    story.append(HRFlowable(width="100%", thickness=2, color=NEON_GREEN))
    story.append(Spacer(1, 0.5 * cm))

    ai = analysis.ai_analysis

    # ── Executive Summary ──────────────────────────────────────────────────────
    story.append(Paragraph("EXECUTIVE SUMMARY", styles["section_title"]))
    story.append(Paragraph(ai.summary, styles["body_text"]))
    story.append(Spacer(1, 0.4 * cm))

    # Who's ahead banner
    ahead_color = NEON_GREEN if "PES" in ai.which_brand_ahead else RED
    story.append(
        Table(
            [[Paragraph(f"⚔️  BATTLEFIELD STATUS: {ai.which_brand_ahead}", styles["banner_text"])]],
            colWidths=["100%"],
            style=TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), DARK_GRAY),
                ("TEXTCOLOR", (0, 0), (-1, -1), ahead_color),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
                ("BOX", (0, 0), (-1, -1), 1.5, ahead_color),
                ("ROUNDEDCORNERS", [4]),
            ]),
        )
    )
    story.append(Spacer(1, 0.4 * cm))

    # ── Scores Comparison ──────────────────────────────────────────────────────
    story.append(Paragraph("COMPETITIVE SCORES", styles["section_title"]))
    scores_data = _build_scores_table(ai.scores)
    if scores_data:
        story.append(scores_data)
    story.append(Spacer(1, 0.4 * cm))

    # ── Daily Battle Report ─────────────────────────────────────────────────
    story.append(Paragraph("TODAY'S BATTLE REPORT", styles["section_title"]))
    story.append(Paragraph(ai.daily_battle_report, styles["body_text"]))
    story.append(Spacer(1, 0.4 * cm))

    # ── SWOT ──────────────────────────────────────────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("SWOT ANALYSIS", styles["section_title"]))
    for brand_key, swot in ai.swot.items():
        story.append(Paragraph(f"▸ {swot.brand}", styles["subsection"]))
        swot_table = _build_swot_table(swot)
        story.append(swot_table)
        story.append(Spacer(1, 0.3 * cm))

    # ── Strategic Recommendations ──────────────────────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("STRATEGIC RECOMMENDATIONS", styles["section_title"]))
    for rec in ai.recommendations:
        priority_color = RED if rec.priority == "high" else YELLOW if rec.priority == "medium" else NEON_GREEN
        story.append(
            Paragraph(
                f'<font color="{priority_color.hexval() if hasattr(priority_color,"hexval") else "#00ff88"}">[{rec.priority.upper()}]</font> {rec.title}',
                styles["rec_title"],
            )
        )
        story.append(Paragraph(rec.description, styles["body_text"]))
        for action in rec.action_items[:4]:
            story.append(Paragraph(f"  → {action}", styles["bullet"]))
        story.append(Paragraph(f"Expected Impact: {rec.expected_impact}", styles["impact"]))
        story.append(Spacer(1, 0.2 * cm))

    # ── AI Suggestions ──────────────────────────────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("AI-GENERATED CAMPAIGN IDEAS", styles["section_title"]))

    story.append(Paragraph("Suggested Ad Hooks", styles["subsection"]))
    for hook in ai.suggested_ad_hooks[:6]:
        story.append(Paragraph(f"  ▸ {hook}", styles["bullet"]))

    story.append(Spacer(1, 0.2 * cm))
    story.append(Paragraph("Reel Ideas", styles["subsection"]))
    for reel in ai.suggested_reel_ideas[:6]:
        story.append(Paragraph(f"  ▸ {reel}", styles["bullet"]))

    story.append(Spacer(1, 0.2 * cm))
    story.append(Paragraph("WhatsApp Campaigns", styles["subsection"]))
    for wc in ai.suggested_whatsapp_campaigns:
        story.append(Paragraph(f"  ▸ {wc}", styles["bullet"]))

    # ── Alerts ────────────────────────────────────────────────────────────────
    if ai.alerts:
        story.append(PageBreak())
        story.append(Paragraph("AI ALERTS", styles["section_title"]))
        for alert in ai.alerts:
            sev = alert.get("severity", "info")
            color = "#ff3b3b" if sev == "critical" else "#ffd60a" if sev == "warning" else "#00d4ff"
            story.append(
                Paragraph(
                    f'<font color="{color}">[{sev.upper()}]</font> {alert.get("title","")} — {alert.get("description","")}',
                    styles["alert_text"],
                )
            )

    # ── Footer ─────────────────────────────────────────────────────────────────
    story.append(Spacer(1, 1 * cm))
    story.append(HRFlowable(width="100%", thickness=1, color=NEON_GREEN))
    story.append(Paragraph("Powered by Jarvis AI — Practical EduSkills War Room", styles["footer"]))

    doc.build(story)
    return buffer.getvalue()


def generate_excel_report(analysis: FullAnalysisResponse) -> bytes:
    wb = openpyxl.Workbook()

    sheets = [
        ("Scores", _excel_scores_sheet),
        ("Ad Analysis", _excel_ads_sheet),
        ("Social Media", _excel_social_sheet),
        ("SWOT", _excel_swot_sheet),
        ("Recommendations", _excel_recommendations_sheet),
        ("Alerts", _excel_alerts_sheet),
    ]

    first = True
    for name, builder in sheets:
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        builder(ws, analysis)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _excel_scores_sheet(ws, analysis: FullAnalysisResponse):
    ai = analysis.ai_analysis
    _style_header(ws, "Competitive Scores Dashboard", cols=4)
    headers = ["Metric", "PES Score", "Nilaya Score", "Leader"]
    _write_row(ws, 3, headers, bold=True, bg="1a1a2e", fg="00ff88")

    metrics = [
        "ad_activity_score", "content_frequency_score", "engagement_score",
        "brand_authority_score", "trust_score", "placement_positioning_score",
        "founder_branding_score", "ai_readiness_score", "innovation_score", "overall_score",
    ]
    pes_s = ai.scores.get("PES")
    nil_s = ai.scores.get("Nilaya")

    for row, metric in enumerate(metrics, start=4):
        label = metric.replace("_", " ").title()
        pes_val = getattr(pes_s, metric, 0) if pes_s else 0
        nil_val = getattr(nil_s, metric, 0) if nil_s else 0
        leader = "PES" if pes_val >= nil_val else "Nilaya"
        _write_row(ws, row, [label, pes_val, nil_val, leader])

    _auto_col_width(ws)


def _excel_ads_sheet(ws, analysis: FullAnalysisResponse):
    _style_header(ws, "Ad Intelligence", cols=5)
    headers = ["Brand", "Active Ads", "Sample Title", "Sample Body", "Platform"]
    _write_row(ws, 3, headers, bold=True, bg="1a1a2e", fg="00d4ff")

    row = 4
    for brand, brand_data in [("PES", analysis.pes_data), ("Nilaya", analysis.nilaya_data)]:
        for ad in brand_data.ads[:10]:
            _write_row(ws, row, [
                brand,
                brand_data.active_ad_count,
                (ad.title or "")[:80],
                (ad.body or "")[:150],
                ad.platform or "",
            ])
            row += 1
    _auto_col_width(ws)


def _excel_social_sheet(ws, analysis: FullAnalysisResponse):
    _style_header(ws, "Social Media Metrics", cols=6)
    headers = ["Brand", "Platform", "Followers", "Posts", "Engagement%", "Posting Freq"]
    _write_row(ws, 3, headers, bold=True, bg="1a1a2e", fg="bf5af2")

    row = 4
    for brand, brand_data in [("PES", analysis.pes_data), ("Nilaya", analysis.nilaya_data)]:
        for s in brand_data.social:
            _write_row(ws, row, [
                brand,
                s.platform.title(),
                s.followers or 0,
                s.posts_count or 0,
                s.engagement_rate or 0,
                s.posting_frequency or "Unknown",
            ])
            row += 1
    _auto_col_width(ws)


def _excel_swot_sheet(ws, analysis: FullAnalysisResponse):
    _style_header(ws, "SWOT Analysis", cols=5)
    headers = ["Brand", "Strengths", "Weaknesses", "Opportunities", "Threats"]
    _write_row(ws, 3, headers, bold=True, bg="1a1a2e", fg="ffd60a")

    row = 4
    for brand_key, swot in analysis.ai_analysis.swot.items():
        _write_row(ws, row, [
            swot.brand,
            "\n".join(f"• {s}" for s in swot.strengths),
            "\n".join(f"• {w}" for w in swot.weaknesses),
            "\n".join(f"• {o}" for o in swot.opportunities),
            "\n".join(f"• {t}" for t in swot.threats),
        ])
        ws.row_dimensions[row].height = max(80, len(swot.strengths) * 15)
        row += 1
    _auto_col_width(ws, max_width=60)


def _excel_recommendations_sheet(ws, analysis: FullAnalysisResponse):
    _style_header(ws, "Strategic Recommendations", cols=5)
    headers = ["Priority", "Category", "Title", "Description", "Expected Impact"]
    _write_row(ws, 3, headers, bold=True, bg="1a1a2e", fg="00ff88")

    for row, rec in enumerate(analysis.ai_analysis.recommendations, start=4):
        _write_row(ws, row, [
            rec.priority.upper(),
            rec.category,
            rec.title,
            rec.description[:200],
            rec.expected_impact,
        ])
    _auto_col_width(ws)


def _excel_alerts_sheet(ws, analysis: FullAnalysisResponse):
    _style_header(ws, "AI Alerts", cols=5)
    headers = ["Severity", "Brand", "Platform", "Title", "Description"]
    _write_row(ws, 3, headers, bold=True, bg="1a1a2e", fg="ff3b3b")

    for row, alert in enumerate(analysis.ai_analysis.alerts, start=4):
        _write_row(ws, row, [
            alert.get("severity", "").upper(),
            alert.get("brand", ""),
            alert.get("platform", ""),
            alert.get("title", ""),
            alert.get("description", "")[:200],
        ])
    _auto_col_width(ws)


# ── PDF helpers ───────────────────────────────────────────────────────────────

def _build_styles():
    base = getSampleStyleSheet()
    styles = {}
    styles["cover_title"] = ParagraphStyle(
        "CoverTitle", fontSize=22, textColor=NEON_GREEN, alignment=TA_CENTER,
        spaceAfter=8, fontName="Helvetica-Bold",
    )
    styles["cover_sub"] = ParagraphStyle(
        "CoverSub", fontSize=14, textColor=NEON_BLUE, alignment=TA_CENTER,
        spaceAfter=4, fontName="Helvetica",
    )
    styles["cover_date"] = ParagraphStyle(
        "CoverDate", fontSize=10, textColor=LIGHT_GRAY, alignment=TA_CENTER,
        spaceAfter=12,
    )
    styles["section_title"] = ParagraphStyle(
        "SectionTitle", fontSize=14, textColor=NEON_GREEN, fontName="Helvetica-Bold",
        spaceBefore=10, spaceAfter=6, borderPad=4,
    )
    styles["subsection"] = ParagraphStyle(
        "Subsection", fontSize=12, textColor=NEON_BLUE, fontName="Helvetica-Bold",
        spaceBefore=6, spaceAfter=4,
    )
    styles["body_text"] = ParagraphStyle(
        "BodyText", fontSize=9, textColor=LIGHT_GRAY, leading=14, spaceAfter=6,
    )
    styles["bullet"] = ParagraphStyle(
        "Bullet", fontSize=9, textColor=LIGHT_GRAY, leftIndent=12, spaceAfter=3,
    )
    styles["banner_text"] = ParagraphStyle(
        "Banner", fontSize=13, textColor=NEON_GREEN, fontName="Helvetica-Bold",
        alignment=TA_CENTER,
    )
    styles["rec_title"] = ParagraphStyle(
        "RecTitle", fontSize=10, textColor=WHITE, fontName="Helvetica-Bold",
        spaceBefore=8, spaceAfter=3,
    )
    styles["impact"] = ParagraphStyle(
        "Impact", fontSize=9, textColor=NEON_PURPLE, fontName="Helvetica-Bold",
        spaceAfter=4,
    )
    styles["alert_text"] = ParagraphStyle(
        "Alert", fontSize=9, textColor=WHITE, spaceAfter=4,
    )
    styles["footer"] = ParagraphStyle(
        "Footer", fontSize=8, textColor=DARK_GRAY, alignment=TA_CENTER, spaceAfter=0,
    )
    return styles


def _build_scores_table(scores: Dict):
    if not scores:
        return None

    metrics = [
        ("Ad Activity", "ad_activity_score"),
        ("Content Frequency", "content_frequency_score"),
        ("Engagement", "engagement_score"),
        ("Brand Authority", "brand_authority_score"),
        ("Trust", "trust_score"),
        ("Placement Positioning", "placement_positioning_score"),
        ("Founder Branding", "founder_branding_score"),
        ("AI Readiness", "ai_readiness_score"),
        ("Innovation", "innovation_score"),
        ("OVERALL", "overall_score"),
    ]

    data = [["Metric", "PES", "Nilaya", "Leader"]]
    for label, key in metrics:
        pes_val = getattr(scores.get("PES"), key, 0) if scores.get("PES") else 0
        nil_val = getattr(scores.get("Nilaya"), key, 0) if scores.get("Nilaya") else 0
        leader = "PES ✓" if pes_val >= nil_val else "Nilaya ✓"
        data.append([label, f"{pes_val:.0f}", f"{nil_val:.0f}", leader])

    table = Table(data, colWidths=["45%", "18%", "18%", "19%"])
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), DARK_GRAY),
            ("TEXTCOLOR", (0, 0), (-1, 0), NEON_GREEN),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("BACKGROUND", (0, 1), (-1, -2), MID_GRAY),
            ("BACKGROUND", (0, -1), (-1, -1), DARK_GRAY),
            ("TEXTCOLOR", (0, 1), (-1, -1), LIGHT_GRAY),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [MID_GRAY, BLACK]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#333355")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )
    return table


def _build_swot_table(swot):
    def _items(lst):
        return "\n".join(f"• {i}" for i in lst[:5]) if lst else "—"

    data = [
        ["STRENGTHS", "WEAKNESSES"],
        [_items(swot.strengths), _items(swot.weaknesses)],
        ["OPPORTUNITIES", "THREATS"],
        [_items(swot.opportunities), _items(swot.threats)],
    ]
    table = Table(data, colWidths=["50%", "50%"])
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#003366")),
            ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#660000")),
            ("BACKGROUND", (0, 2), (0, 2), colors.HexColor("#004400")),
            ("BACKGROUND", (1, 2), (1, 2), colors.HexColor("#664400")),
            ("TEXTCOLOR", (0, 0), (-1, 0), WHITE),
            ("TEXTCOLOR", (0, 2), (-1, 2), WHITE),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("TEXTCOLOR", (0, 1), (-1, 1), LIGHT_GRAY),
            ("TEXTCOLOR", (0, 3), (-1, 3), LIGHT_GRAY),
            ("BACKGROUND", (0, 1), (-1, 1), MID_GRAY),
            ("BACKGROUND", (0, 3), (-1, 3), BLACK),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#333355")),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("ALIGN", (0, 2), (-1, 2), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ])
    )
    return table


# ── Excel helpers ─────────────────────────────────────────────────────────────

def _style_header(ws, title: str, cols: int = 4):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="00FF88")
    ws["A1"].fill = PatternFill(start_color="0A0A1E", end_color="0A0A1E", fill_type="solid")
    ws.merge_cells(f"A1:{get_column_letter(cols)}1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws["A2"] = f"Generated: {datetime.utcnow().strftime('%d %B %Y, %H:%M UTC')}"
    ws["A2"].font = Font(name="Calibri", size=9, color="808080")


def _write_row(ws, row: int, values: list, bold: bool = False, bg: str = None, fg: str = None):
    for col, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=str(val) if not isinstance(val, (int, float)) else val)
        cell.font = Font(
            name="Calibri",
            bold=bold,
            color=fg or "E0E0E0",
        )
        if bg:
            cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        else:
            bg_color = "0D0D2B" if row % 2 == 0 else "16213E"
            cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _auto_col_width(ws, max_width: int = 50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value).split("\n")[0]))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), max_width)
